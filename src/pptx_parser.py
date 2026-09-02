"""PPTX pitch-deck extraction, slide by slide.

Text is walked shape-tree first so grouped shapes and tables inside groups are
not lost, and speaker notes are captured — decks routinely park the pricing or
timeline detail the analysis needs in the notes pane.
"""

from __future__ import annotations

from datetime import date
from pathlib import Path
from typing import Any


def parse(path: str | Path) -> dict[str, Any]:
    """Return {"slides", "doc_date"} for a PPTX deck."""
    from pptx import Presentation

    path = Path(path)
    presentation = Presentation(str(path))
    slides: list[dict[str, Any]] = []

    for number, slide in enumerate(presentation.slides, start=1):
        lines: list[str] = []
        tables: list[list[list[str]]] = []
        for shape in _walk(slide.shapes):
            _collect(shape, lines, tables)

        title = _title(slide) or (lines[0][:120] if lines else "")
        notes = _notes(slide)
        if notes:
            lines.append(f"[speaker notes] {notes}")

        slides.append(
            {
                "index": number,
                "title": title,
                "text": "\n".join(lines).strip(),
                "tables": tables,
            }
        )
    return {"slides": slides, "doc_date": _core_date(presentation)}


def _walk(shapes) -> list[Any]:
    """Flatten the shape tree so grouped shapes are visited too."""
    found: list[Any] = []
    for shape in shapes:
        if getattr(shape, "shape_type", None) is not None and hasattr(shape, "shapes"):
            found.extend(_walk(shape.shapes))
        else:
            found.append(shape)
    return found


def _collect(shape, lines: list[str], tables: list[list[list[str]]]) -> None:
    """Append one shape's text or table content to the slide's accumulators."""
    if getattr(shape, "has_table", False):
        rows = [
            [cell.text.strip() for cell in row.cells]
            for row in shape.table.rows
        ]
        rows = [row for row in rows if any(row)]
        if rows:
            tables.append(rows)
        return

    if getattr(shape, "has_text_frame", False):
        for paragraph in shape.text_frame.paragraphs:
            text = "".join(run.text for run in paragraph.runs).strip()
            if text:
                lines.append(text)
        return

    if getattr(shape, "has_chart", False):
        # Charts carry the numbers a revenue or market slide is built on; the
        # category/series labels are usually enough to read the claim.
        try:
            plot = shape.chart.plots[0]
            categories = [str(category) for category in plot.categories]
            for series in plot.series:
                values = ", ".join("" if v is None else str(v) for v in series.values)
                lines.append(f"[chart] {series.name}: {values}")
            if categories:
                lines.append(f"[chart categories] {', '.join(categories)}")
        except Exception:
            pass


def _title(slide) -> str:
    try:
        if slide.shapes.title is not None:
            return (slide.shapes.title.text or "").strip()[:120]
    except Exception:
        pass
    return ""


def _notes(slide) -> str:
    try:
        if slide.has_notes_slide:
            return (slide.notes_slide.notes_text_frame.text or "").strip()
    except Exception:
        pass
    return ""


def _core_date(presentation) -> date | None:
    """Authored date from the package's core properties."""
    try:
        properties = presentation.core_properties
    except Exception:
        return None
    for stamp in (properties.created, properties.modified):
        if stamp:
            return stamp.date() if hasattr(stamp, "date") else stamp
    return None
